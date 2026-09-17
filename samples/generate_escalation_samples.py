# samples/generate_escalation_samples.py
"""Generate the two hour-scale sample documents.

python-docx lives in the backend image, and docker-compose mounts only ./backend
into it — so the script is fed to the container on stdin, writes into /app, and the
files are copied back out (see the plan's Task 8 Step 2 for the exact commands).

Everything here is fictional. The deadlines are deliberately short — minutes and
hours — because no other sample exercises sub-day timing.
"""

import os
from pathlib import Path

from docx import Document

OUT = Path(os.environ.get("SAMPLES_OUT", Path(__file__).parent))
MARKER = "SAMPLE — fictional document for demonstration"

PROTOCOL = [
    ("Kanangra Court Aged Care — Incident Escalation Protocol", 0),
    (MARKER, None),
    ("Part 1 Purpose", 1),
    ("This protocol sets the escalation steps and timeframes that apply when a "
     "medication incident, fall or unexplained absence occurs at Kanangra Court. "
     "Every provider, resident and reference in this document is fictional.", None),
    ("Part 2 Immediate response", 1),
    ("2.1 Clinical assessment", 2),
    ("The registered nurse on duty must assess the affected resident within 30 "
     "minutes of the incident being discovered and record the assessment in the "
     "clinical record.", None),
    ("2.2 Notifying the nurse in charge", 2),
    ("The registered nurse must notify the nurse in charge within 1 hour of the "
     "incident.", None),
    ("Part 3 Escalation", 1),
    ("3.1 Notifying the family or representative", 2),
    ("The facility manager must notify the resident's nominated representative "
     "within 4 hours of the incident.", None),
    ("3.2 Notifying the prescriber", 2),
    ("Where a medication was given in error, the prescriber must be contacted "
     "within 2 hours of the error being identified.", None),
    ("3.3 Reportable incident notification", 2),
    ("A Priority 1 reportable incident must be notified to the Aged Care Quality "
     "and Safety Commission within 24 hours of the provider becoming aware of it.", None),
    ("Part 4 Review", 1),
    ("4.1 Preliminary review", 2),
    ("The clinical governance lead must complete a preliminary review within 2 "
     "business days of the incident.", None),
    ("4.2 Full investigation", 2),
    ("A full investigation report must be submitted to the board within 30 days "
     "and 4 hours of the incident, to align with the Commission's reporting cycle.", None),
    ("4.3 Ongoing monitoring", 2),
    ("Medication administration audits must continue monthly for 6 months after "
     "any Priority 1 medication incident.", None),
]

CASE_STUDY = [
    ("Kanangra Court Aged Care — Case study: medication error, Resident K", 0),
    (MARKER, None),
    ("Incident summary", 1),
    ("The medication error occurred on 14 September 2026 at 3:10pm. Resident K was "
     "administered a dose of a medication prescribed for another resident during the "
     "afternoon medication round in Wing B.", None),
    ("Immediate response taken", 1),
    ("The registered nurse assessed Resident K at 3:25pm. Observations were within "
     "normal limits and were recorded in the clinical record at 3:40pm.", None),
    ("The nurse in charge was notified at 3:35pm and attended the wing.", None),
    ("The resident's daughter, the nominated representative, was telephoned at "
     "4:15pm and informed of the error and the observations taken.", None),
    ("Outstanding at the time of writing", 1),
    ("The prescriber had not been contacted at the time this record was written. "
     "The clinical governance lead has been asked to begin the preliminary review "
     "but has not yet scheduled it.", None),
    ("No notification has been made to the Aged Care Quality and Safety Commission; "
     "the facility manager is seeking advice on whether the incident meets the "
     "Priority 1 threshold.", None),
]


def _write(blocks, path: Path) -> None:
    doc = Document()
    for text, level in blocks:
        if level is None:
            doc.add_paragraph(text)
        else:
            doc.add_heading(text, level=level)
    doc.save(path)
    print(f"wrote {path}")


if __name__ == "__main__":
    _write(PROTOCOL, OUT / "Kanangra-Court-Incident-Escalation-Protocol.docx")
    _write(CASE_STUDY, OUT / "Kanangra-Court-Case-Study-Medication-Error.docx")
